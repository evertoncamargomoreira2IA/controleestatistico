"""Gera o método de uso a partir do padrão Foxconn e dos prints reais do QMS."""
from pathlib import Path
from copy import deepcopy
import argparse
import hashlib
import json
import zipfile
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--template', type=Path, default=ROOT/'docs/manual/modelo-foxconn.docx')
args = parser.parse_args()
reference = args.template.resolve()
doc = Document(reference)
source_bytes = reference.read_bytes()
def preserve_package_parts(path):
    editable = {'word/document.xml','word/_rels/document.xml.rels','word/header1.xml','word/footer1.xml','word/settings.xml','word/styles.xml','docProps/core.xml','[Content_Types].xml'}
    original = zipfile.ZipFile(io.BytesIO(source_bytes))
    generated = zipfile.ZipFile(io.BytesIO(path.read_bytes()))
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for part in generated.infolist():
            data = original.read(part.filename) if part.filename in original.namelist() and part.filename not in editable else generated.read(part.filename)
            target.writestr(part, data)
    path.write_bytes(output.getvalue())
TITLE = 'Método de uso da coleta CEP no QMS'
BLUE = '164A85'
DATE = '17/10/2026'
OUT = ROOT/'docs/manual/Metodo de Uso QMS CEP Sprint 1.docx'
WORK = ROOT/'artifacts/manual'
WORK.mkdir(parents=True, exist_ok=True)
PRINTS = ROOT/'docs/manual/prints'

# Keep the reference's section geometry, logo, furniture and style package.
section = doc.sections[0]
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)
for row in section.header.tables[0].rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                if 'Instrução para Operação do Fluid Dispenser HP INKJET' in run.text:
                    run.text = run.text.replace('Instrução para Operação do Fluid Dispenser HP INKJET', TITLE)
                if 'FBR-EN-2-662' in run.text:
                    run.text = run.text.replace('FBR-EN-2-662', 'MINUTA QMS CEP')
for paragraph in section.footer.paragraphs:
    for run in paragraph.runs:
        if 'FBR-EN-2-662' in run.text:
            run.text = run.text.replace('FBR-EN-2-662', 'QMS CEP — MINUTA')
settings = doc.settings.element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields'); settings.append(update)
update.set(qn('w:val'), 'true')

# A body-free reusable template makes future sprint manuals independent of the desktop reference.
template_path = ROOT/'docs/manual/modelo-foxconn.docx'
if not template_path.exists() or reference != template_path.resolve():
    doc.save(template_path)
    preserve_package_parts(template_path)

def text(value, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_together = True
    r = p.add_run(value); r.font.name = 'Arial'; r.font.size = Pt(9); r.bold = bold
    return p

def heading(value):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(value); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    return p

def cell_text(cell, value, bold=False, color='000000', size=8):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1
    r = p.add_run(str(value)); r.font.name = 'Arial'; r.font.size = Pt(size); r.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, data, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement('w:'+edge)
        for key,value in [('val','single'),('sz','5'),('color','8793A0')]: b.set(qn('w:'+key), value)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for col,width in zip(t.columns,widths): col.width = Inches(width)
    for idx,h in enumerate(headers):
        cell=t.rows[0].cells[idx]; cell.width=Inches(widths[idx]); cell_text(cell,h,True,'FFFFFF')
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),BLUE); cell._tc.get_or_add_tcPr().append(shd)
    header=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    for values in data:
        row=t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for idx,value in enumerate(values): row.cells[idx].width=Inches(widths[idx]); cell_text(row.cells[idx],value)
    text('')
    return t

def figure(filename, caption, crop=None, width=7.02):
    path=PRINTS/filename
    if crop:
        out=WORK/('crop-'+filename)
        with Image.open(path) as image: image.crop(crop).save(out)
        path=out
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path),width=Inches(width))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7)
    r=p.add_run(caption); r.font.name='Arial'; r.font.size=Pt(8); r.italic=True; r.font.color.rgb=RGBColor.from_string('5C6670')

def page():
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(1)

if 'Title' not in doc.styles:
    doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after=Pt(8)
r=p.add_run(TITLE); r.font.name='Arial'; r.font.size=Pt(15); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
table(['Controle','Identificação','Controle','Identificação'],[
    ['Código','A definir pelo controle de documentos','Revisão','00 — minuta'],
    ['Data',DATE,'Situação','Em validação — Sprint 1'],
    ['Elaboração','Engenharia de Processos','Revisão','A preencher'],
    ['Aprovação','A preencher','Aplicação','QMS CEP web'],
], [1,2.55,1,2.53])
heading('1 FINALIDADE')
text('Orientar o registro de conjuntos CEP de 25 amostras com três leituras cada, a seleção de dados existentes e a consulta do progresso e do histórico. Esta revisão descreve as telas entregues na Sprint 1, com exemplos de demonstração.')
heading('2 APLICAÇÃO')
text('Aplicável a técnicos, Engenharia e Qualidade que utilizam o QMS pelo navegador. A coleta completa reúne 75 leituras. Nesta sprint são apresentados média e amplitude; cartas de controle e análise de capacidade pertencem às próximas entregas.')
heading('3 DEFINIÇÕES')
table(['Termo','Uso no sistema'],[
    ['Amostra ou subgrupo','Uma linha da grade, com três leituras da mesma característica.'],
    ['Média','Soma das três leituras dividida por três; calculada apenas com a linha completa.'],
    ['Amplitude','Maior leitura menos a menor leitura da mesma amostra.'],
    ['Rascunho','Conjunto salvo que ainda aceita edição, inclusive com linhas incompletas.'],
    ['Concluído','Conjunto com 25 amostras completas, preservado em consulta. Não significa aprovação estatística.'],
], [1.35,5.73])
heading('4 RESPONSABILIDADES E PRÉ REQUISITOS')
text('Técnico ou administrador: conferir o rastreio, registrar e revisar as leituras. Engenharia: manter cadastros, revisão e especificações. Leitor: consultar conjuntos e histórico sem alterar os registros.')
text('Antes de iniciar, verificar se a característica está ativa, configurada como X̄-R e com subgrupo de três leituras. Usar instrumento ativo e com calibração válida na data da coleta. Acessar o endereço fornecido pelo responsável pelo sistema.')
text('Ambiente local de demonstração: usuário admin e senha admin. A configuração está limitada a Development; o acesso corporativo será definido na homologação. Documento sem aprovação formal nesta revisão.',True)

page(); heading('5 CRIAR E IDENTIFICAR O CONJUNTO')
text('5.1 Acessar Conjuntos CEP 25 × 3 no menu ou Novo conjunto CEP no painel. Selecionar a característica pelo caminho Área / Produto / Linha / Posto / Tooling e informar um título que identifique a coleta.')
figure('01-conjunto.png','Figura 1 — Seleção do processo e criação do rascunho.',(260,0,1340,700))
text('5.2 Clicar em Criar rascunho. O sistema gera o número do conjunto, registra a revisão da característica e preserva nominal, LSL e USL usados nessa coleta.')
text('5.3 Para digitação ou colagem, informar instrumento, data e hora local, turno e lote. Clicar em Aplicar identificação às linhas digitadas. Quando os horários ou instrumentos forem diferentes, expandir Frontend na linha e ajustar o rastreio individual.')
text('5.4 Não misturar características ou revisões no mesmo conjunto. Se a característica deixar de aparecer na lista, conferir se está ativa, X̄-R e com subgrupo de três leituras. Os cadastros de teste da seção 11 já atendem a essa configuração.')

page(); heading('6 DIGITAR OU COLAR AS LEITURAS')
text('6.1 Digitar as três medições em cada linha da grade. É permitido usar vírgula ou ponto decimal. Conferir média, amplitude e identificação da amostra; o sistema mantém as leituras originais.')
figure('03-grade.png','Figura 2 — Grade com cinco amostras completas e suas médias.',(270,410,1320,890))
text('6.2 Para colar do Excel, expandir Colar leituras do Excel e colar uma linha por amostra, com três colunas. Usar TAB, ponto e vírgula ou espaço como separador. Não incluir cabeçalho; o limite é 25 linhas.')
text('Exemplo de linha: 10,02;10,04;10,03. A média correspondente é 10,03 e a amplitude é 0,02. Valores de exemplo não substituem medições reais.')
text('6.3 Marcar a confirmação de substituição da grade e clicar em Conferir e aplicar colagem. A grade anterior será substituída; conferir os valores e salvar. Uma colagem inválida não deve ser corrigida removendo dados sem rastreio.')
text('6.4 Clicar em Salvar rascunho antes de sair. Confirmar a mensagem de salvamento. O indicador da grade pode refletir alterações ainda não salvas; a mensagem identifica quando elas existem.')
figure('02-rascunho.png','Figura 3 — Progresso de preenchimento de um rascunho reaberto.',(270,0,1320,350))

page(); heading('7 IMPORTAR ARQUIVO E ESCOLHER DADOS DA BASE')
text('7.1 Para CSV ou XLSX, usar o template da característica disponível em Coleta. Preencher as três leituras e os campos de rastreio. Acessar Importar CSV ou XLSX na base, selecionar o arquivo e conferir a prévia.')
figure('04-importacao.png','Figura 4 — Prévia do arquivo demonstrativo com 25 linhas válidas.',(270,0,1320,640))
text('7.2 Corrigir as linhas com erro antes de prosseguir. A confirmação importa somente as linhas válidas. Conferir a quantidade efetivamente importada. O limite é 5 MB e 5.000 linhas; reimportar o mesmo conteúdo é bloqueado.')
text('7.3 Criar ou abrir o conjunto CEP e expandir Selecionar amostras da base. Informar datas, turno ou lote e clicar em Buscar na base. A busca mostra registros da mesma característica e revisão, com três leituras e não excluídos.')
text('7.4 Marcar explicitamente até 25 amostras. Conferir instrumento, operador, origem e horário. Se houver mais de 500 resultados, reduzir o período ou aplicar outros filtros.')

page(); heading('8 VINCULAR E CONCLUIR O CONJUNTO')
figure('05-base.png','Figura 5 — Seleção explícita das 25 amostras da base.',(270,130,1320,880))
text('8.1 Marcar a confirmação de substituição da grade e clicar em Aplicar selecionadas à grade. As amostras são ordenadas por data e ID. Seus valores e rastreio ficam vinculados aos registros originais, sem edição das leituras nessa grade.')
text('8.2 Clicar em Salvar rascunho. Para concluir, conferir 25 / 25 amostras e 75 / 75 leituras, verificar o rastreio, marcar a confirmação de conferência e clicar em Concluir conjunto.')
text('8.3 Após a conclusão, a grade fica somente para consulta. Para uma nova coleta ou correção posterior, criar outro conjunto e manter a referência do anterior. A conclusão não libera produto nem confirma estabilidade do processo.',True)

page(); heading('9 ACOMPANHAR O PROGRESSO E CONSULTAR O HISTÓRICO')
text('9.1 Acessar Progresso diário CEP. Selecionar a data de criação e clicar em Consultar data. Usar o filtro para localizar processo, título ou característica.')
figure('08-progresso.png','Figura 6 — Painel diário com um rascunho e um conjunto concluído.',(260,0,1340,850))
text('9.2 O painel mostra o estado atual dos conjuntos criados na data consultada, no horário de São Paulo. O percentual é de preenchimento: cinco amostras completas equivalem a 20%; 25 equivalem a 100%.')
text('9.3 Esta revisão ainda não mede cumprimento de programação ou meta do turno. O quadro de programação e as pendências diárias serão evoluídos na Sprint 3.')

page(); heading('10 HISTÓRICO E TRATAMENTO DE IMPEDIMENTOS')
text('10.1 Clicar em Abrir e consultar no painel diário. Ao final da grade, consultar Histórico do conjunto. Cada evento informa horário local, usuário e operação; Consultar detalhes apresenta os registros antes e depois do salvamento.')
figure('07-historico.png','Figura 7 — Eventos de criação, salvamento e conclusão do mesmo conjunto.',(270,530,1320,880))
table(['Situação','Conduta do usuário'],[
    ['Conjunto incompleto','Salvar rascunho e continuar a coleta. Concluir exige 25 linhas com três leituras cada.'],
    ['Número inválido','Corrigir a posição indicada. Não usar texto, NaN ou infinito.'],
    ['Instrumento ou data inválida','Conferir instrumento ativo, calibração, data e hora real. Não alterar datas para contornar uma rejeição.'],
    ['Conjunto atualizado por outra pessoa','Anotar a divergência e reabrir o conjunto antes de editar. Não sobrescrever o trabalho do outro usuário.'],
    ['Característica alterada ou inativada','Preservar o conjunto anterior e criar outro com a revisão atual.'],
    ['Amostra da base incompatível ou excluída','Repetir a consulta e conferir a origem; não reutilizar valores alterados manualmente como se fossem da base.'],
], [2.0,5.08])
text('O histórico operacional do conjunto está disponível aos perfis de consulta. A tela geral Histórico tem permissões próprias de auditoria. A retenção dos registros e os backups corporativos serão definidos com a TI e a Qualidade.')

page(); heading('11 ESPECIFICAÇÕES CARREGADAS PARA TESTE')
text('Os 11 dispositivos abaixo foram transcritos da imagem fornecida pelo solicitante. Os códigos AT e PK e os níveis Linha de teste e Posto de validação são identificadores provisórios de demonstração. Não foram inventadas medições para esses dispositivos.')
table(['Código e produto','Descrição','Nominal mm','LSL mm','USL mm'],[
    ['AT-001 Athena Lebi','Feed Shaft Gear Insert Press Fixture','326,95','326,65','327,25'],
    ['AT-002 Athena Lebi','Feed Shaft Pulley Insert Press Fixture','4,68','4,58','4,78'],
    ['AT-003 Athena Lebi','Belt Tension - Paper Motor','−5,25','−6,00','−4,50'],
    ['AT-004 Athena Lebi','X-Bias Press Fixture','49,39','49,29','49,49'],
    ['AT-005 Athena Lebi','Out Put Clutch Press Fixture','1,11','1,01','1,21'],
    ['PK-001 Peak Mech Sayan','Gear Feedroller 24T Press','64,13','63,93','64,33'],
    ['PK-002 Peak Mech Sayan','Pulley Cluster and Encoder Ring Press','4,73','4,33','4,83'],
    ['PK-003 Peak Mech Sayan','Gear Right Feedroller 17T Press','317,28','317,18','317,58'],
    ['PK-004 Peak Mech Sayan','Output Gear 42T Press','10,09','9,89','10,29'],
    ['PK-005 Peak Mech Sayan','Pick Encoder Disk and Hub Press','40,45','40,15','40,75'],
    ['PK-006 Peak Mech Sayan','Gear Feedroller 24T Pick & Press','48,63','48,43','48,83'],
], [1.15,3.25,.9,.9,.88])
heading('12 PRÓXIMAS REVISÕES E CONTROLE DO DOCUMENTO')
text('A imagem de Control Chart fornecida pelo solicitante é a referência da próxima tela de análise: medições por média, distribuição lateral, nominal, média global, limites de controle e especificação distintos e indicadores de regras. Essa tela ainda não integra a Sprint 1.')
text('Os critérios visuais solicitados são Cp acima de 1,67, Cpk acima de 1,33 e sequências de sete pontos crescentes acima do nominal ou decrescentes abaixo do nominal. O método estatístico e os limites serão implementados e testados nas sprints correspondentes.')
table(['Rev.','Data','Descrição da revisão','Revisão e aprovação'],[['00',DATE,'Emissão da minuta do método da Sprint 1 com prints reais de demonstração, histórico, progresso e especificações de teste.','A preencher']], [.4,.8,4.45,1.43])

doc.core_properties.title=TITLE
doc.core_properties.subject='Método operacional da Sprint 1 do QMS CEP'
doc.core_properties.author='Engenharia de Processos'
doc.save(OUT)
preserve_package_parts(OUT)
assert hashlib.sha256(reference.read_bytes()).digest()==hashlib.sha256(source_bytes).digest(), 'Referência alterada'
print(OUT)
