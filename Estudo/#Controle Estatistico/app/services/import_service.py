import csv
from datetime import datetime, timezone
import hashlib
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.collection import ImportBatch
from app.models.masterdata import Characteristic, Instrument, Tooling
from app.services.audit import record
from app.services.collection_service import save_subgroup, validate_subgroup
from app.services.timezone_service import LOCAL_TZ

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 5000
HEADERS = [
    "tooling_code",
    "characteristic_name",
    "characteristic_id",
    "data_hora",
    "instrumento_codigo",
    "turno",
    "lote",
    "observacao",
]


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_upload(filename, content: bytes) -> list[dict]:
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("O arquivo excede o limite de 5 MB.")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        try:
            sheet = load_workbook(BytesIO(content), data_only=True, read_only=True)["Coleta"]
        except (BadZipFile, KeyError, ValueError, OSError) as exc:
            raise ValueError("Arquivo XLSX inválido ou sem a aba Coleta.") from exc
        values = sheet.iter_rows(values_only=True)
        rows = list(values)
    elif suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        rows = list(csv.reader(StringIO(text), delimiter=delimiter))
    else:
        raise ValueError("Envie um arquivo .xlsx ou .csv.")
    if not rows:
        return []
    header = [str(value or "").strip().lower() for value in rows[0]]
    index = {name: position for position, name in enumerate(header)}
    if not all(name in index for name in HEADERS):
        raise ValueError("O arquivo não contém os cabeçalhos obrigatórios do template.")
    value_headers = sorted(
        (name for name in header if name.startswith("v") and name[1:].isdigit()),
        key=lambda name: int(name[1:]),
    )
    if not value_headers:
        raise ValueError("O arquivo precisa conter colunas v1, v2, ...")
    parsed = []
    for line_number, raw in enumerate(rows[1:], start=2):
        if not any(not _blank(value) for value in raw):
            continue
        item = {name: raw[index[name]] if index[name] < len(raw) else None for name in HEADERS}
        item["values"] = [
            raw[index[name]] if index[name] < len(raw) else None for name in value_headers
        ]
        if str(item.get("observacao") or "").strip().lower() == "exemplo":
            continue
        item["linha"] = line_number
        parsed.append(item)
        if len(parsed) > MAX_ROWS:
            raise ValueError("O arquivo excede o limite de 5000 linhas de dados.")
    return parsed


def _parse_timestamp(value):
    if isinstance(value, datetime):
        parsed = value
    else:
        text = str(value or "").strip()
        try:
            parsed = datetime.strptime(text, "%d/%m/%Y %H:%M")
        except ValueError:
            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            except ValueError as exc:
                raise ValueError("Data e hora inválidas; use dd/mm/aaaa hh:mm ou ISO.") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=LOCAL_TZ)
    return parsed.astimezone(timezone.utc)


def preview(db: Session, rows: list[dict], user) -> list[dict]:
    output = []
    for position, row in enumerate(rows, start=2):
        line = row.get("linha", position)
        errors = []
        characteristic = None
        instrument = None
        timestamp = None
        tooling_code = str(row.get("tooling_code") or "").strip()
        try:
            ident = int(row.get("characteristic_id"))
            characteristic = db.get(Characteristic, ident)
        except (TypeError, ValueError):
            errors.append("ID da característica inválido.")
        if not characteristic:
            errors.append("Característica não encontrada.")
        else:
            tooling = db.get(Tooling, characteristic.tooling_id)
            if row.get("characteristic_name") != characteristic.name:
                errors.append("O nome da característica não corresponde ao ID.")
            if not tooling or tooling.code != tooling_code:
                errors.append("A característica não corresponde ao tooling informado.")
        instrument_code = str(row.get("instrumento_codigo") or "").strip()
        instrument = db.query(Instrument).filter_by(code=instrument_code).first()
        if not instrument:
            errors.append("Instrumento não encontrado pelo código informado.")
        try:
            timestamp = _parse_timestamp(row.get("data_hora"))
        except (TypeError, ValueError) as exc:
            errors.append(str(exc))
        shift = str(row.get("turno") or "").strip()
        if shift not in {"T1", "T2", "T3"}:
            errors.append("Turno inválido; informe T1, T2 ou T3.")
        values = row.get("values", [])
        if characteristic:
            check_time = timestamp or datetime.now(timezone.utc)
            parsed, validation_errors = validate_subgroup(
                characteristic, values, instrument, check_time
            )
            errors.extend(validation_errors)
            if not validation_errors:
                # Store the parsed floats so commit never re-reads "9,8"-style text.
                values = parsed
        output.append(
            {
                "linha": line,
                "characteristic": characteristic,
                "instrument": instrument,
                "collected_at": timestamp,
                "values": values,
                "errors": errors,
                "shift": shift,
                "lot": row.get("lote"),
                "note": row.get("observacao"),
            }
        )
    return output


def commit_import(
    db: Session, filename, content: bytes, rows_preview: list[dict], user
) -> ImportBatch:
    digest = hashlib.sha256(content).hexdigest()
    previous = db.query(ImportBatch).filter_by(file_sha256=digest).first()
    if previous:
        uploaded_at = previous.uploaded_at
        if uploaded_at.tzinfo is None:
            uploaded_at = uploaded_at.replace(tzinfo=timezone.utc)
        raise ValueError(
            f"Este arquivo já foi importado em {uploaded_at.astimezone(LOCAL_TZ):%d/%m/%Y} "
            f"por {previous.uploader.full_name}."
        )
    valid = [row for row in rows_preview if not row["errors"]]
    rejected = len(rows_preview) - len(valid)
    batch = ImportBatch(
        filename=Path(filename).name[:255],
        uploaded_by=user.id,
        uploaded_at=datetime.now(timezone.utc),
        rows_total=len(rows_preview),
        rows_imported=len(valid),
        rows_rejected=rejected,
        file_sha256=digest,
        status="completed" if not rejected else "completed_with_errors",
    )
    db.add(batch)
    db.flush()
    for row in valid:
        save_subgroup(
            db,
            row["characteristic"],
            row["values"],
            user,
            row["instrument"],
            row["collected_at"],
            row["shift"],
            str(row.get("lot") or "").strip() or None,
            "upload",
            str(row.get("note") or "").strip() or None,
            batch.id,
        )
    record(
        db,
        actor_id=user.id,
        entity="import_batch",
        entity_id=batch.id,
        action="import.commit",
        after={
            "filename": batch.filename,
            "rows_total": batch.rows_total,
            "rows_imported": batch.rows_imported,
            "rows_rejected": batch.rows_rejected,
            "file_sha256": digest,
        },
        origin="web",
    )
    db.commit()
    return batch
