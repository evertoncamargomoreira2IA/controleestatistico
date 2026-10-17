import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

from app.models.collection import ImportBatch, Subgroup
from app.models.masterdata import Characteristic, Instrument
from tests.api.test_auth_users import csrf, db_for, seed, signin
from tests.api.test_collection import create_instrument, setup_characteristic


def csv_content(characteristic_id, tooling_code, instrument_code):
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        [
            "tooling_code",
            "characteristic_name",
            "characteristic_id",
            "data_hora",
            "instrumento_codigo",
            "turno",
            "lote",
            "observacao",
            "v1",
            "v2",
            "v3",
        ]
    )
    for values in (("9,8", "10", "10.2"), ("9.7", "10", "10.1"), ("bad", "10", "10")):
        writer.writerow(
            [
                tooling_code,
                "Medida Upload",
                characteristic_id,
                "15/10/2020 12:00",
                instrument_code,
                "T1",
                "L-1",
                "",
                *values,
            ]
        )
    return output.getvalue().encode("utf-8-sig")


def post_preview(client, content, filename="dados.csv"):
    return client.post(
        "/coleta/importar/previa",
        data={"csrf_token": csrf(client)},
        files={"file": (filename, content)},
    )


def confirm(client, token, filename="dados.csv"):
    return client.post(
        "/coleta/importar/confirmar",
        data={"csrf_token": csrf(client), "token": token, "filename": filename},
    )


def test_template_and_csv_preview_confirm_and_duplicate(client):
    seed(client)
    signin(client)
    characteristic_id = setup_characteristic(client, "Medida Upload")
    instrument_id = create_instrument(client, "CAL-UPLOAD")
    db = db_for(client)
    from app.models.masterdata import Tooling

    characteristic = db.get(Characteristic, characteristic_id)
    tooling_code = db.get(Tooling, characteristic.tooling_id).code
    instrument_code = db.get(Instrument, instrument_id).code
    db.close()

    response = client.get(f"/coleta/template.xlsx?characteristic_id={characteristic_id}")
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert response.status_code == 200
    assert workbook["Coleta"].cell(1, 9).value == "v1"
    assert workbook["Coleta"].cell(1, 11).value == "v3"

    content = csv_content(characteristic_id, tooling_code, instrument_code)
    response = post_preview(client, content)
    assert response.status_code == 200
    assert "2 linhas válidas" in response.text
    assert "1 linhas com erro" in response.text
    token = response.text.split('name="token" value="')[1].split('"')[0]
    response = confirm(client, token)
    assert response.status_code == 200
    db = db_for(client)
    batch = db.query(ImportBatch).one()
    assert (batch.rows_imported, batch.rows_rejected) == (2, 1)
    assert db.query(Subgroup).count() == 2
    assert all(row.source == "upload" for row in db.query(Subgroup).all())
    db.close()

    response = post_preview(client, content)
    token = response.text.split('name="token" value="')[1].split('"')[0]
    response = confirm(client, token)
    assert response.status_code == 400
    assert "já foi importado" in response.text


def test_template_xlsx_round_trip_and_invalid_uuid(client):
    seed(client)
    signin(client)
    characteristic_id = setup_characteristic(client, "XLSX Upload")
    instrument_id = create_instrument(client, "CAL-XLSX")
    db = db_for(client)
    instrument_code = db.get(Instrument, instrument_id).code
    db.close()

    response = client.get(f"/coleta/template.xlsx?characteristic_id={characteristic_id}")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Coleta"]
    sheet["D2"] = "15/10/2020 12:00"
    sheet["E2"] = instrument_code
    sheet["F2"] = "T1"
    sheet["H2"] = ""
    sheet["I2"], sheet["J2"], sheet["K2"] = 9.8, 10, 10.2
    buffer = BytesIO()
    workbook.save(buffer)
    preview_response = post_preview(client, buffer.getvalue(), "coleta.xlsx")
    assert preview_response.status_code == 200
    token = preview_response.text.split('name="token" value="')[1].split('"')[0]
    assert confirm(client, token, "coleta.xlsx").status_code == 200
    assert (
        client.post(
            "/coleta/importar/confirmar",
            data={"csrf_token": csrf(client), "token": "invalid", "filename": "coleta.xlsx"},
        ).status_code
        == 400
    )


def test_reader_cannot_open_import_form(client):
    seed(client, "leitor", "reader")
    signin(client, "reader")
    assert client.get("/coleta/importar").status_code == 403
